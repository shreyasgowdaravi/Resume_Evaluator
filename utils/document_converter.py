# Copyright (c) 2025 GradientM IT Consulting & Services Pvt Ltd
# All rights reserved.

# Fix for hashlib.md5 compatibility with xhtml2pdf
import hashlib
original_md5 = hashlib.md5
def patched_md5(*args, **kwargs):
    kwargs.pop('usedforsecurity', None)
    return original_md5(*args, **kwargs)
hashlib.md5 = patched_md5

import os
import tempfile
from docx import Document
from xhtml2pdf import pisa
from pdf2docx import Converter
try:
    import PyPDF2
    from PyPDF2 import PdfReader
except ImportError:
    try:
        import pypdf as PyPDF2
        from pypdf import PdfReader
    except ImportError:
        PyPDF2 = None
        PdfReader = None
from io import BytesIO
from werkzeug.utils import secure_filename

def docx_to_pdf(docx_input):
    """Convert DOCX file to PDF format"""
    try:
        # Handle both file objects and file paths
        if isinstance(docx_input, str):
            if not os.path.exists(docx_input):
                raise Exception(f"DOCX file not found: {docx_input}")
            doc = Document(docx_input)
        else:
            # Handle file-like objects (uploaded files)
            if hasattr(docx_input, 'seek'):
                docx_input.seek(0)  # Reset file pointer
            
            # Save to temporary file first
            temp_docx = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
            if hasattr(docx_input, 'read'):
                temp_docx.write(docx_input.read())
            else:
                # If it's already bytes
                temp_docx.write(docx_input)
            temp_docx.close()
            
            try:
                doc = Document(temp_docx.name)
            finally:
                # Clean up temporary file
                try:
                    os.unlink(temp_docx.name)
                except:
                    pass
        
        # Extract text and basic formatting
        html_content = """
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <style>
                body { font-family: Arial, sans-serif; margin: 20px; line-height: 1.6; }
                p { margin-bottom: 10px; }
                .bold { font-weight: bold; }
                .heading { font-size: 16px; font-weight: bold; color: #333; margin-top: 15px; }
            </style>
        </head>
        <body>
        """
        
        paragraph_count = 0
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                paragraph_count += 1
                # Escape HTML characters
                import html
                escaped_text = html.escape(paragraph.text.strip())
                
                # Check if paragraph has bold formatting or is a heading
                is_bold = any(run.bold for run in paragraph.runs if run.bold)
                is_heading = paragraph.style.name.startswith('Heading') if paragraph.style else False
                
                if is_heading or (is_bold and len(escaped_text) < 100):
                    html_content += f'<p class="heading">{escaped_text}</p>\n'
                elif is_bold:
                    html_content += f'<p class="bold">{escaped_text}</p>\n'
                else:
                    html_content += f'<p>{escaped_text}</p>\n'
        
        # Also extract text from tables if any
        for table in doc.tables:
            html_content += '<table style="border-collapse: collapse; width: 100%; margin: 10px 0;">\n'
            for row in table.rows:
                html_content += '<tr>\n'
                for cell in row.cells:
                    if cell.text.strip():
                        escaped_cell = html.escape(cell.text.strip())
                        html_content += f'<td style="border: 1px solid #ddd; padding: 5px;">{escaped_cell}</td>\n'
                html_content += '</tr>\n'
            html_content += '</table>\n'
        
        if paragraph_count == 0:
            raise Exception("No text content found in DOCX file")
        
        html_content += "</body></html>"
        
        # Create temporary PDF file
        temp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        temp_pdf.close()
        
        # Convert HTML to PDF with better settings
        with open(temp_pdf.name, "wb") as pdf_file:
            pisa_status = pisa.CreatePDF(
                html_content.encode('utf-8'),
                dest=pdf_file,
                encoding='utf-8',
                show_error_as_pdf=True
            )
            
        if pisa_status.err:
            print(f"PDF conversion warning: {pisa_status.err}")
            # Don't raise exception for warnings, only for critical errors
        
        # Verify the PDF was created and has content
        if not os.path.exists(temp_pdf.name) or os.path.getsize(temp_pdf.name) == 0:
            raise Exception("Generated PDF is empty or was not created")
            
        return temp_pdf.name
        
    except Exception as e:
        raise Exception(f"DOCX to PDF conversion failed: {str(e)}")

def pdf_to_docx(pdf_file):
    """Convert PDF file to DOCX format"""
    try:
        # Save uploaded file to temporary location
        temp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        pdf_file.save(temp_pdf.name)
        
        # Create temporary DOCX file
        temp_docx = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
        temp_docx.close()
        
        # Convert PDF to DOCX
        cv = Converter(temp_pdf.name)
        cv.convert(temp_docx.name)
        cv.close()
        
        # Clean up temporary PDF
        os.unlink(temp_pdf.name)
        
        return temp_docx.name
        
    except Exception as e:
        # Fallback: extract text using PyPDF2 and create simple DOCX
        try:
            if PyPDF2 is None:
                raise Exception("PyPDF2 not available")
                
            pdf_file.seek(0)  # Reset file pointer
            pdf_reader = PdfReader(pdf_file)
            text_content = ""
            
            for page in pdf_reader.pages:
                text_content += page.extract_text() + "\n"
            
            # Create DOCX with extracted text
            doc = Document()
            for line in text_content.split('\n'):
                if line.strip():
                    doc.add_paragraph(line.strip())
            
            temp_docx = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
            doc.save(temp_docx.name)
            temp_docx.close()
            
            return temp_docx.name
            
        except Exception as fallback_error:
            raise Exception(f"PDF to DOCX conversion failed: {str(e)}. Fallback also failed: {str(fallback_error)}")

def extract_text_from_file(file_path):
    """Extract text from various file formats"""
    try:
        if not os.path.exists(file_path):
            raise Exception(f"File not found: {file_path}")
        
        if file_path.endswith(".txt"):
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
                if not content.strip():
                    raise Exception("TXT file is empty")
                return content
        elif file_path.endswith(".docx"):
            try:
                doc = Document(file_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                if not paragraphs:
                    raise Exception("No text content found in DOCX file")
                return "\n".join(paragraphs)
            except Exception as docx_error:
                raise Exception(f"Failed to read DOCX file: {str(docx_error)}")
        elif file_path.endswith(".pdf"):
            if PyPDF2 is None or PdfReader is None:
                raise Exception("PyPDF2/pypdf not available for PDF text extraction")
            with open(file_path, 'rb') as f:
                pdf_reader = PdfReader(f)
                text_content = ""
                for page in pdf_reader.pages:
                    text_content += page.extract_text() + "\n"
                if not text_content.strip():
                    raise Exception("No text content found in PDF file")
                return text_content
        else:
            raise Exception(f"Unsupported file format: {file_path}")
    except Exception as e:
        raise Exception(f"Text extraction failed: {str(e)}")

def text_to_pdf(text_content, title="Document"):
    """Convert plain text to PDF format"""
    try:
        if not text_content or not text_content.strip():
            raise Exception("Text content is empty")
        
        # Escape HTML characters and preserve line breaks
        import html
        escaped_content = html.escape(text_content.strip())
        formatted_content = escaped_content.replace('\n', '</p><p>')
        
        html_content = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <style>
                body {{ font-family: Arial, sans-serif; margin: 20px; line-height: 1.6; }}
                h1 {{ color: #333; border-bottom: 2px solid #667eea; padding-bottom: 10px; }}
                p {{ margin-bottom: 10px; }}
            </style>
        </head>
        <body>
            <h1>{html.escape(title)}</h1>
            <p>{formatted_content}</p>
        </body>
        </html>
        """
        
        temp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        temp_pdf.close()
        
        with open(temp_pdf.name, "wb") as pdf_file:
            pisa_status = pisa.CreatePDF(
                html_content.encode('utf-8'),
                dest=pdf_file,
                encoding='utf-8'
            )
            
        if pisa_status.err:
            raise Exception(f"PDF conversion error: {pisa_status.err}")
        
        # Verify the PDF was created
        if not os.path.exists(temp_pdf.name) or os.path.getsize(temp_pdf.name) == 0:
            raise Exception("Generated PDF is empty or was not created")
            
        return temp_pdf.name
        
    except Exception as e:
        raise Exception(f"Text to PDF conversion failed: {str(e)}")

def get_file_extension(filename):
    """Get file extension in lowercase"""
    return os.path.splitext(filename)[1].lower()

def is_supported_resume_format(filename):
    """Check if resume file format is supported"""
    supported_formats = ['.pdf', '.docx']
    return get_file_extension(filename) in supported_formats

def is_supported_jd_format(filename):
    """Check if job description file format is supported"""
    supported_formats = ['.txt', '.docx', '.pdf']
    return get_file_extension(filename) in supported_formats