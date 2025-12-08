#!/usr/bin/env python3
"""
Test script to verify DOCX to PDF conversion and PDF merging fixes
"""

import os
import tempfile
from utils.document_converter import docx_to_pdf, text_to_pdf
from utils.pdf_merger import merge_resume_and_report, merge_with_jd_and_report

def create_test_docx():
    """Create a simple test DOCX file"""
    from docx import Document
    
    doc = Document()
    doc.add_heading('Test Resume', 0)
    doc.add_paragraph('Name: John Doe')
    doc.add_paragraph('Email: john.doe@example.com')
    doc.add_paragraph('Experience: 5 years in software development')
    doc.add_paragraph('Skills: Python, JavaScript, SQL')
    
    temp_docx = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
    doc.save(temp_docx.name)
    temp_docx.close()
    
    return temp_docx.name

def create_test_report_pdf():
    """Create a simple test report PDF"""
    report_content = """
    Test Evaluation Report
    
    Candidate: John Doe
    Score: 85/100
    Verdict: Shortlist
    
    This is a test evaluation report.
    """
    
    return text_to_pdf(report_content, "Test Evaluation Report")

def test_docx_to_pdf_conversion():
    """Test DOCX to PDF conversion"""
    print("Testing DOCX to PDF conversion...")
    
    try:
        # Create test DOCX
        docx_path = create_test_docx()
        print(f"Created test DOCX: {docx_path}")
        
        # Convert to PDF
        pdf_path = docx_to_pdf(docx_path)
        print(f"Converted to PDF: {pdf_path}")
        
        # Verify PDF exists and has content
        if os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0:
            print("✓ DOCX to PDF conversion successful!")
            
            # Cleanup
            os.unlink(docx_path)
            os.unlink(pdf_path)
            return True
        else:
            print("X PDF conversion failed - file empty or doesn't exist")
            return False
            
    except Exception as e:
        print(f"X DOCX to PDF conversion failed: {e}")
        return False

def test_pdf_merging():
    """Test PDF merging with DOCX resume"""
    print("\nTesting PDF merging with DOCX resume...")
    
    try:
        # Create test files
        docx_path = create_test_docx()
        report_path = create_test_report_pdf()
        
        print(f"Created test DOCX resume: {docx_path}")
        print(f"Created test report PDF: {report_path}")
        
        # Test merging
        merged_path = merge_resume_and_report(docx_path, report_path, "test_merged.pdf")
        
        if merged_path and os.path.exists(merged_path) and os.path.getsize(merged_path) > 0:
            print("✓ PDF merging with DOCX resume successful!")
            print(f"Merged PDF created: {merged_path}")
            
            # Cleanup
            os.unlink(docx_path)
            os.unlink(report_path)
            os.unlink(merged_path)
            return True
        else:
            print("X PDF merging failed")
            return False
            
    except Exception as e:
        print(f"X PDF merging failed: {e}")
        return False

def test_complete_package_merging():
    """Test complete package merging (JD + Report + Resume)"""
    print("\nTesting complete package merging...")
    
    try:
        # Create test files
        docx_path = create_test_docx()
        report_path = create_test_report_pdf()
        
        # Create JD text file
        jd_content = """
        Job Description: Software Developer
        
        Requirements:
        - 3+ years experience in Python
        - Knowledge of web frameworks
        - Database experience
        """
        
        jd_path = tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False, encoding='utf-8')
        jd_path.write(jd_content)
        jd_path.close()
        jd_path = jd_path.name
        
        print(f"Created test JD: {jd_path}")
        print(f"Created test DOCX resume: {docx_path}")
        print(f"Created test report PDF: {report_path}")
        
        # Test complete merging
        merged_path = merge_with_jd_and_report(jd_path, docx_path, report_path, "test_complete.pdf")
        
        if merged_path and os.path.exists(merged_path) and os.path.getsize(merged_path) > 0:
            print("✓ Complete package merging successful!")
            print(f"Complete package PDF created: {merged_path}")
            
            # Cleanup
            os.unlink(docx_path)
            os.unlink(report_path)
            os.unlink(jd_path)
            os.unlink(merged_path)
            return True
        else:
            print("X Complete package merging failed")
            return False
            
    except Exception as e:
        print(f"X Complete package merging failed: {e}")
        return False

def main():
    """Run all tests"""
    print("Testing DOCX and PDF merging fixes...\n")
    
    results = []
    
    # Test 1: DOCX to PDF conversion
    results.append(test_docx_to_pdf_conversion())
    
    # Test 2: PDF merging with DOCX resume
    results.append(test_pdf_merging())
    
    # Test 3: Complete package merging
    results.append(test_complete_package_merging())
    
    # Summary
    print(f"\nTest Results:")
    print(f"Passed: {sum(results)}/{len(results)}")
    print(f"Failed: {len(results) - sum(results)}/{len(results)}")
    
    if all(results):
        print("\nAll tests passed! DOCX and PDF merging fixes are working correctly.")
    else:
        print("\nSome tests failed. Please check the error messages above.")

if __name__ == "__main__":
    main()